"""add_market_slide.py - close the two open deck gaps before submission.

Two things, both recorded in docs/submission/:

1. COMPLIANCE finding 10. lablab's Presentation rubric only reaches its top two
   bands with market analysis, and the 10-page deck had no market slide. This
   inserts one as slide 10, cloned from slide 9's four-stat layout, carrying the
   exact numbers already on screen in the demo video's market card
   (docs/submission/_closing/close_market.html), so the deck and the video
   cannot contradict each other. Sources are in DEMO_NOTES.md, "The closing
   cards".

2. TRUTH_AUDIT row 22. Slide 8 said the LLM "only picks which clause applies".
   That is true of Gemini (claim_check.py returns a verdict + clause_id and
   nothing else) but NOT of the AssemblyAI Voice Agent, whose model emits the
   spoken utterance from a "say exactly" instruction (server/voice.py:56,467)
   and is graded only after the audio has gone out (server/voice.py:448-450).
   The card body and the slide footnote are rewritten to be true of both legs.

The deck grows 10 -> 11 pages, so every "NN / 10" marker is renumbered.

    python tools/add_market_slide.py            # rewrites the pptx in place
    python tools/add_market_slide.py --check    # reports, writes nothing

Idempotent: a second run finds the market slide already present and the slide 8
text already rewritten, and exits 0 without changing anything. Re-export
ClauseCatcher.pdf afterwards (PowerPoint COM; see the bottom of this file).
"""
from __future__ import annotations

import argparse
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "docs/submission/ClauseCatcher.pptx"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"

MARKET_KICKER = "THE MARKET, SIZED BOTTOM-UP"

# (shape name on the cloned slide, new text) - shape names come straight from
# slide 9, which this slide is a copy of.
MARKET_TEXT = {
    "Text 0": MARKET_KICKER,
    "Text 2": "The category, and the slice we can name",
    # four stat cards: big number + grey caption underneath
    "Text 4": "$1.6B–$32B",
    "Text 5": "2026 conversation-intelligence estimates",
    "Text 7": "1.59M",
    "Text 8": "US wholesale & manufacturing reps (BLS 2025)",
    "Text 10": "$0.6B–$1.1B",
    "Text 11": "those reps at a $30–60 per-seat hypothesis",
    "Text 13": "292k",
    "Text 14": "technical & scientific reps: the beachhead",
    # bullets
    "Text 16": "Each firm draws the category differently — the spread is the honest "
               "number, not the top of it",
    "Text 18": "Beachhead: the 292,000 technical and scientific reps, who sell the "
               "hardest contracts",
    "Text 20": "$30–60 per seat per month is our hypothesis — no buyer has been "
               "asked to pay it yet",
    # footnote
    "Text 21": "1.59M = 1.3M except-technical-and-scientific + 292k technical and "
               "scientific, BLS Occupational Outlook Handbook 2025 employment "
               "(bls.gov/ooh). $0.6B–$1.1B = 1.59M × $30–60 × 12.",
}

# The video card's strings are kept character for character, and "$0.6B-$1.1B"
# is far too wide for slide 9's 40 pt stat box - it broke mid-word at 32 pt in a
# render. Drop the stat to 28 pt and let the box run to the card's inner edge.
STAT_SHAPES = ("Text 4", "Text 7", "Text 10", "Text 13")
STAT_PT = 28
STAT_WIDTH = Emu(2156384)  # card inner width from the stat's own left edge

# (exact current paragraph text, replacement, reason)
SLIDE8_EDITS = [
    ("Everything spoken or shown comes from the uploaded contract. The LLM only "
     "picks which clause applies.",
     # keep this <= ~105 chars: the card body wraps at ~27 and a 5th line
     # renders flush against the card's bottom edge (checked in a PNG export)
     "The words are the contract's own. Gemini picks the clause ID; the Voice "
     "Agent is told to say it verbatim.",
     "TRUTH_AUDIT #22 - 'the LLM only picks the clause' is false of the Voice Agent leg"),
    ("A false accusation is worse than a missed one.",
     "A false accusation is worse than a missed one. The read-back is graded "
     "against the clause after it is spoken: drift is flagged, not blocked.",
     "TRUTH_AUDIT #22 - voice.py:448-450 grades after the audio is already out"),
]


def set_text(para, new: str) -> None:
    """Replace a paragraph's text, keeping the first run's formatting."""
    para.runs[0].text = new
    for run in para.runs[1:]:
        run.text = ""


def only_para(shape):
    return next(p for p in shape.text_frame.paragraphs if "".join(r.text for r in p.runs).strip())


def clone_slide(prs, src, index: int):
    """Copy a slide's shape tree onto a new slide and move it to `index`.

    python-pptx has no slide copy. Slide 9 is plain text boxes and autoshapes
    with no picture or chart parts, so a deepcopy of the spTree children needs
    no relationship fixing.
    """
    new = prs.slides.add_slide(src.slide_layout)
    for shape in list(new.shapes):  # drop anything the layout seeded
        shape._element.getparent().remove(shape._element)
    for shape in src.shapes:
        new.shapes._spTree.append(copy.deepcopy(shape._element))

    # The dark navy is a <p:bg> on the slide's own cSld, not on the layout.
    # Without this the clone renders white-on-white and the copy vanishes.
    src_bg = src._element.find(f"{{{NS_P}}}cSld/{{{NS_P}}}bg")
    if src_bg is not None:
        new_csld = new._element.find(f"{{{NS_P}}}cSld")
        new_csld.insert(0, copy.deepcopy(src_bg))

    sld_id_lst = prs.slides._sldIdLst
    entry = list(sld_id_lst)[-1]
    sld_id_lst.remove(entry)
    sld_id_lst.insert(index, entry)
    return new


def has_market_slide(prs) -> bool:
    return any(
        MARKET_KICKER in shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="report only, do not write")
    ap.add_argument("--deck", default=str(DECK))
    args = ap.parse_args()

    prs = Presentation(args.deck)
    changed = 0

    # --- 1. the market slide -------------------------------------------------
    if has_market_slide(prs):
        print("market slide already present - skipping")
    else:
        src = prs.slides[8]  # slide 9, "MEASURED, NOT PROJECTED"
        new = clone_slide(prs, src, 9)  # lands as slide 10, before "what's next"
        by_name = {s.name: s for s in new.shapes}
        missing = sorted(set(MARKET_TEXT) - set(by_name))
        if missing:
            sys.exit(f"slide 9 shape names changed; not found: {missing}")
        for name, text in MARKET_TEXT.items():
            set_text(only_para(by_name[name]), text)
        for name in STAT_SHAPES:
            shape = by_name[name]
            shape.width = STAT_WIDTH
            only_para(shape).runs[0].font.size = Pt(STAT_PT)
        print(f"added market slide as slide 10 ({len(prs.slides)} slides total)")
        changed += 1

    # --- 2. slide 8 safety wording ------------------------------------------
    total = len(prs.slides)
    for slide_no, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs).strip()
                if not text:
                    continue
                for old, new, why in SLIDE8_EDITS:
                    if text == new:
                        break
                    if text == old:
                        print(f"slide {slide_no}: rewrote {text[:40]!r}...  ({why})")
                        if not args.check:
                            set_text(para, new)
                        changed += 1
                        break
                else:
                    # --- 3. page markers, including the new slide ------------
                    if text.endswith("/ 10") and text[:2].isdigit():
                        marker = f"{slide_no:02d} / {total}"
                        if marker != text:
                            print(f"slide {slide_no}: {text!r} -> {marker!r}")
                            if not args.check:
                                set_text(para, marker)
                            changed += 1

    if args.check:
        print(f"\ncheck only: {changed} change(s) pending")
        sys.exit(1 if changed else 0)
    if not changed:
        print("nothing to do")
        return
    prs.save(args.deck)
    print(f"\nwrote {args.deck} - {changed} change(s)")
    print("Now re-export the PDF:")
    print(r"""  powershell -NoProfile -Command "$app=New-Object -ComObject PowerPoint.Application; """
          r"""$pres=$app.Presentations.Open('<repo>\docs\submission\ClauseCatcher.pptx',$true,$false,$false); """
          r"""$pres.SaveAs('<repo>\docs\submission\ClauseCatcher.pdf', 32); $pres.Close(); $app.Quit()" """)


if __name__ == "__main__":
    main()
